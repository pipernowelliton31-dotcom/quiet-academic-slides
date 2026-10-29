#!/usr/bin/env python3
"""Audit key structural/typographic risks in an editable PPTX.

Usage:
  python audit_deck.py deck.pptx [--json report.json] [--strict]

This scanner is a guardrail, not a visual renderer: every slide still needs
pixel-level visual QA after rendering. No external font files are bundled.
"""
import argparse
import io
import json
import re
import zipfile
from collections import Counter
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU_PER_INCH = 914400
BAD_GLYPHS = ('\ufffd', '\u25a1', '\u25af', '\u0378')
SUSPICIOUS = re.compile(r'(?:Ã.|Â[\x80-\xbf]|Ð.|Ø[\x80-\xbf])')


def walk(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(shape.shapes)


def image_ar(shape):
    try:
        with Image.open(io.BytesIO(shape.image.blob)) as im:
            return im.width / im.height
    except Exception:
        return None


def is_cropped(shape):
    try:
        # Native image cropping is intentional if a nonzero source rectangle exists.
        return any(abs(getattr(shape, key, 0) or 0) > 0.0001
                   for key in ('crop_top','crop_bottom','crop_left','crop_right'))
    except Exception:
        return False


def audit(path):
    errors, warnings = [], []
    fonts, font_sizes, fill_colors = Counter(), Counter(), Counter()
    with zipfile.ZipFile(path) as z:
        bad_zip = z.testzip()
        if bad_zip:
            errors.append(f'ZIP integrity failure at {bad_zip}')
    pres = Presentation(path)
    w, h = pres.slide_width, pres.slide_height
    for index, slide in enumerate(pres.slides, start=1):
        for sh in walk(slide.shapes):
            # Tiny EMU tolerance prevents rounding false positives.
            x0, y0, x1, y1 = sh.left, sh.top, sh.left + sh.width, sh.top + sh.height
            if x0 < -5000 or y0 < -5000 or x1 > w + 5000 or y1 > h + 5000:
                errors.append(f'Slide {index}: object outside canvas: {sh.name} ({x0/EMU_PER_INCH:.2f},{y0/EMU_PER_INCH:.2f})')
            try:
                if sh.fill.type and sh.fill.fore_color.type and sh.fill.fore_color.rgb:
                    fill_colors[str(sh.fill.fore_color.rgb)] += 1
            except Exception:
                pass
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                iar = image_ar(sh)
                sar = (sh.width / sh.height) if sh.height else None
                if iar and sar and not is_cropped(sh) and abs((sar / iar) - 1) > 0.12:
                    warnings.append(f'Slide {index}: image {sh.name} aspect mismatch {iar:.2f} source vs {sar:.2f} canvas; inspect crop/stretch')
            if getattr(sh,'has_text_frame',False):
                text = sh.text
                if any(c in text for c in BAD_GLYPHS) or SUSPICIOUS.search(text):
                    errors.append(f'Slide {index}: possible garbled text in {sh.name}: {text[:80]!r}')
                for paragraph in sh.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if not run.text.strip():
                            continue
                        name = run.font.name or '(theme/inherited)'
                        fonts[name] += 1
                        if run.font.size:
                            pts = round(run.font.size.pt, 1)
                            font_sizes[pts] += 1
                            if pts < 8:
                                warnings.append(f'Slide {index}: very small text {pts}pt: {run.text[:45]!r}')
    # Ensure a flag for the type problem that originally broke the reference deck.
    variable_risks = [name for name in fonts if any(t in name.lower() for t in ('variable','inter variable','noto serif display'))]
    if variable_risks:
        warnings.append(f'Variable or historically problematic fonts found: {variable_risks}')
    return {
        'deck': str(path), 'slides': len(pres.slides),
        'size_inches': [round(w/EMU_PER_INCH, 3), round(h/EMU_PER_INCH,3)],
        'fonts': fonts.most_common(), 'font_sizes': font_sizes.most_common(),
        'fill_colors': fill_colors.most_common(),
        'errors': errors, 'warnings': warnings,
        'status': 'FAIL' if errors else 'REVIEW' if warnings else 'PASS'
    }


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('pptx',type=Path)
    p.add_argument('--json', dest='json_out',type=Path)
    p.add_argument('--strict',action='store_true',help='Exit nonzero on structural errors')
    args=p.parse_args()
    if not args.pptx.exists():
        p.error('PPTX not found')
    report=audit(args.pptx)
    print(f"Deck: {args.pptx.name}; slides={report['slides']}; result={report['status']}")
    print('Fonts:',', '.join(f'{n} ({k})' for n,k in report['fonts']))
    for item in report['errors']:
        print('ERROR:',item)
    for item in report['warnings'][:30]:
        print('REVIEW:',item)
    if len(report['warnings'])>30:
        print(f"... {len(report['warnings'])-30} additional review warnings in report JSON")
    if args.json_out:
        args.json_out.parent.mkdir(parents=True,exist_ok=True)
        args.json_out.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    if args.strict and report['errors']:
        return 1
    return 0


if __name__=='__main__':
    raise SystemExit(main())
